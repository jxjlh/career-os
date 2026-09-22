"""简历文件解析器.

支持 PDF / Word / 纯文本 / 图片简历，采用「本地抽取优先、视觉 OCR 兜底」的混合降级策略：
1. PDF 用 pdfplumber 抽文字；Word 用 python-docx；txt/md 直接解码。
2. 若本地抽取结果过短（扫描版 PDF / 无文本层），自动降级为视觉模型 OCR。
3. 图片简历直接走视觉 OCR。
"""
from __future__ import annotations

import base64
import io
from typing import Any

import httpx

from app.core.config import get_settings
from app.core.errors import AppError

_MIN_TEXT_CHARS = 80  # 本地抽取低于此字数视为扫描件，触发 OCR
_OCR_TIMEOUT = httpx.Timeout(connect=15.0, read=120.0, write=30.0, pool=30.0)
_MAX_IMAGE_EDGE = 1800
_JPEG_QUALITY = 88

_IMAGE_EXTS = {"png", "jpg", "jpeg", "webp", "bmp", "tif", "tiff", "heic"}
_TEXT_EXTS = {"txt", "md", "markdown", "text"}


def _extension(filename: str) -> str:
    return (filename.rsplit(".", 1)[-1] if "." in filename else "").lower()


def _decode_text(content: bytes) -> str:
    for enc in ("utf-8", "gb18030", "utf-16"):
        try:
            return content.decode(enc)
        except UnicodeDecodeError:
            continue
    return content.decode("utf-8", errors="replace")


def _extract_pdf_text(content: bytes) -> str:
    import pdfplumber

    parts: list[str] = []
    with pdfplumber.open(io.BytesIO(content)) as pdf:
        for page in pdf.pages:
            text = page.extract_text() or ""
            if text:
                parts.append(text)
    return "\n".join(parts).strip()


def _extract_docx_text(content: bytes) -> str:
    import docx

    document = docx.Document(io.BytesIO(content))
    parts: list[str] = [p.text for p in document.paragraphs if p.text.strip()]
    for table in document.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells if c.text.strip()]
            if cells:
                parts.append(" | ".join(cells))
    return "\n".join(parts).strip()


def _render_pdf_to_images(content: bytes) -> list[bytes]:
    """用 pypdfium2 把 PDF 每页渲染成 PNG，供视觉 OCR 使用."""
    import pypdfium2 as pdfium

    pdf = pdfium.PdfDocument(content)
    images: list[bytes] = []
    try:
        for i in range(len(pdf)):
            page = pdf[i]
            bitmap = page.render(scale=2.0)
            pil = bitmap.to_pil()
            buf = io.BytesIO()
            pil.save(buf, format="PNG")
            images.append(buf.getvalue())
    finally:
        pdf.close()
    return images


def _preprocess_image(image: bytes) -> bytes:
    from PIL import Image as PILImage

    try:
        img = PILImage.open(io.BytesIO(image))
        if img.mode in ("RGBA", "P", "LA"):
            img = img.convert("RGB")
        w, h = img.size
        longest = max(w, h)
        if longest > _MAX_IMAGE_EDGE:
            ratio = _MAX_IMAGE_EDGE / longest
            img = img.resize((int(w * ratio), int(h * ratio)), PILImage.LANCZOS)
        buf = io.BytesIO()
        img.save(buf, format="JPEG", quality=_JPEG_QUALITY, optimize=True)
        return buf.getvalue()
    except Exception:
        return image


async def _vision_ocr_images(images: list[bytes]) -> str:
    settings = get_settings()
    if not settings.openai_api_key:
        raise AppError(
            code="RESUME_OCR_NOT_CONFIGURED",
            message="这份简历没有可提取的文字层，且当前未配置视觉识别能力。请改用文字版 PDF，或直接在文本框粘贴简历内容。",
            status=503,
        )
    model = settings.finance_ocr_model or "gpt-4o-mini"
    parts: list[str] = []
    async with httpx.AsyncClient(timeout=_OCR_TIMEOUT) as client:
        for idx, image in enumerate(images):
            img = _preprocess_image(image)
            data_url = f"data:image/jpeg;base64,{base64.b64encode(img).decode('ascii')}"
            payload: dict[str, Any] = {
                "model": model,
                "temperature": 0,
                "messages": [
                    {
                        "role": "user",
                        "content": [
                            {
                                "type": "text",
                                "text": (
                                    "这是一份简历的扫描图片（第 %d 页）。请逐字转录图中全部文字，"
                                    "保持原有的顺序、段落与换行，不要概括、不要补写、不要翻译，"
                                    "只输出图片上真实存在的文字。" % (idx + 1)
                                ),
                            },
                            {"type": "image_url", "image_url": {"url": data_url, "detail": "high"}},
                        ],
                    }
                ],
            }
            resp = await client.post(
                f"{settings.openai_base_url.rstrip('/')}/chat/completions",
                headers={"Authorization": f"Bearer {settings.openai_api_key}"},
                json=payload,
            )
            resp.raise_for_status()
            result = resp.json()
            text = result["choices"][0]["message"]["content"]
            if isinstance(text, str) and text.strip():
                parts.append(text.strip())
    if not parts:
        raise AppError(
            code="RESUME_OCR_EMPTY",
            message="视觉识别没能从这份简历里读取出文字，请改传文字版 PDF 或直接粘贴简历内容。",
            status=422,
        )
    return "\n".join(parts)


async def extract_resume_text(content: bytes, filename: str, content_type: str) -> tuple[str, str]:
    """解析简历文件，返回 (text, method)."""
    ext = _extension(filename)
    is_image = (content_type or "").startswith("image/") or ext in _IMAGE_EXTS

    if is_image:
        text = await _vision_ocr_images([content])
        return text, "vision_ocr"

    if ext == "pdf":
        text = _extract_pdf_text(content)
        if len(text) >= _MIN_TEXT_CHARS:
            return text, "pdf_text"
        images = _render_pdf_to_images(content)
        if images:
            text = await _vision_ocr_images(images)
            return text, "vision_ocr"
        raise AppError(
            code="RESUME_PDF_UNPARSEABLE",
            message="这份 PDF 无法解析，请上传文字版 PDF 或直接粘贴简历内容。",
            status=422,
        )

    if ext == "docx":
        text = _extract_docx_text(content)
        if text.strip():
            return text, "docx_text"
        raise AppError(code="RESUME_DOCX_EMPTY", message="这份 Word 文档没有可提取的文字。", status=422)

    if ext in _TEXT_EXTS:
        text = _decode_text(content).strip()
        if text:
            return text, "txt"
        raise AppError(code="RESUME_TEXT_EMPTY", message="文本文件内容为空。", status=422)

    # 未知类型：先尝试按文本解码，再按图片 OCR
    text = _decode_text(content).strip()
    if text and len(text) >= _MIN_TEXT_CHARS:
        return text, "txt"
    try:
        return await _vision_ocr_images([content]), "vision_ocr"
    except AppError:
        raise AppError(
            code="RESUME_UNSUPPORTED_FORMAT",
            message="不支持的简历格式。请上传 PDF / Word / 图片 / 纯文本。",
            status=422,
        )
